import pytest
from openpyxl import load_workbook
from backend.core.models import Video
from backend.services import export_excel as excel


def test_headers_order_fills_hyperlinks_and_summary(tmp_path):
    old = Video('old', 'https://example/old', 'a', page_name='A', views=6600000,
                transcript='old speech', taken_on='2026-10-01')
    new = Video('new', 'https://example/new', 'b', page_name='B', views=66000,
                transcript='new speech', so_kenh=2, kenh_cung_dang=['B', 'C'])
    line = 'B (facebook): quét 2 video'
    path = excel.export([old, new], [line], ['Bỏ qua video', 'không đọc được kênh'], tmp_path)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ['Nội dung', 'Tình hình quét']
    sheet = workbook['Nội dung']
    assert [cell.value for cell in sheet[1]] == excel.HEADERS
    assert sheet['C2'].value == 'new speech' and sheet['C3'].value == 'old speech'
    assert sheet['J2'].value == 'Mới' and sheet['J3'].value == 'Đã lấy ngày 01/10'
    assert sheet['D2'].value is None and sheet['B2'].value is None
    assert sheet['A2'].fill.fgColor.rgb.endswith('E2EFDA')
    assert sheet['J2'].fill.fgColor.rgb.endswith('E2EFDA')
    assert sheet['A3'].fill.patternType is None
    assert sheet['A1'].fill.fgColor.rgb.endswith('1F4E78') and sheet['A1'].font.bold
    assert sheet['E2'].hyperlink.target == new.url
    assert sheet['I2'].value == '2 (B, C)' and sheet.freeze_panes == 'A2'
    assert sheet.column_dimensions['C'].width == 90
    workbook.close()
    result = excel.read_summary(path)
    assert (result['total'], result['new'], result['taken_before']) == (2, 1, 1)
    assert result['views_total'] == 6666000
    assert result['top'][0]['url'] == old.url
    assert result['scan_lines'] == [line]
    assert result['notes'] == ['Bỏ qua video']
    assert result['errors'] == ['LỖI: không đọc được kênh']


def test_collision_and_formula_safety(tmp_path):
    video = Video('1', 'https://example/1', 'a', page_name='@name', transcript=' =SUM(A1:A2)')
    first = excel.export([video], output_dir=tmp_path)
    second = excel.export([video], output_dir=tmp_path)
    assert first != second and '(lần 2)' in second.name and first.exists()
    workbook = load_workbook(first)
    assert workbook['Nội dung']['C2'].value.startswith("'")
    assert workbook['Nội dung']['F2'].value == "'@name"
    workbook.close()


@pytest.mark.parametrize('value,expected', [(6600000, '6,6 triệu'), (1000000, '1,0 triệu'),
                                           (66000, '66 nghìn'), (165000, '165 nghìn'), (12, '12'), (None, '')])
def test_view_format(value, expected):
    assert excel.format_views(value) == expected


def test_invalid_workbook(tmp_path):
    path = tmp_path / 'invalid.xlsx'
    path.write_text('not a workbook')
    with pytest.raises(ValueError, match='Excel'):
        excel.read_summary(path)
