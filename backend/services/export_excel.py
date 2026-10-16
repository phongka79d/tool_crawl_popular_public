"""Styled public reports and backward-compatible library summaries."""
from __future__ import annotations
from collections import Counter
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from backend.core import config
from backend.core.models import Video
from backend.core.normalize import parse_count, to_vn_datetime, VN_TZ

SHEET_DETAIL = 'Chi tiết video'
SHEET_SUMMARY = 'Tổng hợp kênh'
SHEET_LOG = 'Nhật ký'
HEADERS = ['STT', 'Kênh', 'Follower kênh', 'Ngày đăng', 'Định dạng', 'Thời lượng (giây)', 'Link bài',
           'Lượt xem', 'Cảm xúc', 'Bình luận', 'Chia sẻ', 'Tổng tương tác', 'View/Follower',
           'Hệ số so trung vị kênh', 'Tỷ lệ tương tác', 'Tỷ lệ share',
           'Chủ đề', 'Caption', 'Nội dung thoại', 'Hook gốc',
           'Số kênh cùng đăng', 'Tình trạng', 'Trạng thái', 'Ghi chú']
LEGACY_HEADERS = ['STT', 'Chủ đề', 'Nội dung thoại', 'Trạng thái', 'Link video mẫu', 'Kênh',
                  'Lượt xem', 'Ngày đăng', 'Số kênh cùng đăng', 'Tình trạng']
WIDTHS = [7, 24, 14, 17, 14, 13, 16, 14, 14, 14, 14, 14, 16, 18, 16, 14,
          28, 60, 60, 60, 24, 22, 16, 60]
GROUPS = [(1, 7, 'Thông tin bài', 'DCEAF3'), (8, 16, 'Chỉ số', 'DDF0E8'),
          (17, 20, 'Nội dung', 'E6EFF3'), (21, 24, 'Theo dõi', 'E8F2EF')]
FMT_DATETIME = 'dd/mm/yyyy hh:mm'
FMT_DEC2 = '#,##0.00'
FMT_PCT1 = '0.0%'
FMT_PCT2 = '0.00%'
HEADER_FILL = PatternFill('solid', fgColor='234E63')
HEADER_FONT = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
TITLE_FONT = Font(name='Calibri', bold=True, color='234E63', size=14)
BODY_FONT = Font(name='Calibri', color='243746', size=11)
ZEBRA_FILL = PatternFill('solid', fgColor='F3F7FA')
TABLE_BORDER = Border(**{edge: Side(style='thin', color='BFC7CE')
                         for edge in ('left', 'right', 'top', 'bottom')})
TOP_WRAP = Alignment(vertical='top', wrap_text=True)
HEADER_ALIGN = Alignment(vertical='center', wrap_text=True)


def format_views(value: int | float | None) -> str:
    if value is None:
        return ''
    if value >= 1000000:
        return f'{value / 1000000:.1f}'.replace('.', ',') + ' triệu'
    if value >= 1000:
        return f'{round(value / 1000)} nghìn'
    return str(int(value)) if value == int(value) else str(value).replace('.', ',')


def _safe(value):
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub('', value)
    if value.lstrip().startswith(('=', '+', '-', '@')):
        value = "'" + value
    return value[:32767]


def _put(sheet, row, column, value, fmt=None):
    cell = sheet.cell(row, column, _safe(value))
    cell.font = BODY_FONT
    cell.alignment = TOP_WRAP
    if fmt:
        cell.number_format = fmt
    return cell


def first_lines_hook(caption: str, transcript: str) -> str:
    """Use the first caption line (200 chars) and first 30 raw transcript words."""
    lines = (caption or '').strip().splitlines()
    first = lines[0].strip()[:200].strip() if lines else ''
    spoken = ' '.join((transcript or '').split()[:30])
    return '\n'.join(part for part in (first, spoken) if part)


def _status(video):
    return ('Đã lấy ngày ' + datetime.strptime(video.taken_on, '%Y-%m-%d').strftime('%d/%m')
            if video.taken_on else 'Mới')


def _details(sheet, videos):
    for first, last, title, colour in GROUPS:
        sheet.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)
        for column in range(first, last + 1):
            cell = sheet.cell(1, column)
            cell.fill = PatternFill('solid', fgColor=colour)
            cell.border = TABLE_BORDER
        cell = _put(sheet, 1, first, title)
        cell.font = TITLE_FONT
        cell.alignment = HEADER_ALIGN
    for column, (name, width) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = _put(sheet, 2, column, name)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, TABLE_BORDER
        cell.alignment = HEADER_ALIGN
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 44
    for index, video in enumerate(videos, 1):
        channels = str(video.so_kenh)
        if video.so_kenh > 1:
            channels += ' (' + ', '.join(video.kenh_cung_dang) + ')'
        notes = list(video.notes)
        if video.transcript_clean:
            notes.append('Lời thoại đã được AI sửa')
        elif video.transcript:
            notes.append('Bản chép máy, chưa sửa')
        posted = to_vn_datetime(video.posted_at).replace(tzinfo=None) if video.posted_at is not None else None
        values = [index, video.page_name or video.page_url, format_views(video.page_followers), posted,
                  {'video': 'Video thường', 'reel': 'Reel'}.get(video.format, video.format), video.duration_s,
                  video.url, *[format_views(value) for value in
                  (video.views, video.reactions, video.comments, video.shares, video.interactions)],
                  video.ratio_follow, video.ratio_median, video.engagement, video.share_rate,
                  '', video.caption, video.transcript_clean or video.transcript,
                  first_lines_hook(video.caption, video.transcript), channels, _status(video), '',
                  '\n'.join(dict.fromkeys(notes))]
        row = index + 2
        formats = {4: FMT_DATETIME, 6: '#,##0', 13: FMT_PCT1, 14: FMT_DEC2,
                   15: FMT_PCT1, 16: FMT_PCT2}
        for column, value in enumerate(values, 1):
            cell = _put(sheet, row, column, value, formats.get(column))
            cell.border = TABLE_BORDER
            if row % 2:
                cell.fill = ZEBRA_FILL
        if video.url.startswith(('https://', 'http://')):
            cell = sheet.cell(row, 7)
            cell.hyperlink = video.url
            cell.hyperlink.display = 'Mở bài ↗'
            cell.number_format = ';;;"Mở bài ↗"'
            cell.font = Font(name='Calibri', color='0563C1', underline='single', size=11)
        if not video.taken_on:
            for column in (1, 22):
                sheet.cell(row, column).fill = PatternFill('solid', fgColor='E2EFDA')
        sheet.row_dimensions[row].height = 90
    sheet.freeze_panes = 'C3'
    sheet.auto_filter.ref = f'A2:X{len(videos) + 2}'
    sheet.print_title_rows = '1:2'
    sheet.sheet_view.zoomScale = 85


def _table(sheet, start, title, headers, rows):
    _put(sheet, start, 1, title).font = TITLE_FONT
    sheet.merge_cells(start_row=start, start_column=1, end_row=start, end_column=len(headers))
    sheet.row_dimensions[start].height = 28
    for column, header in enumerate(headers, 1):
        cell = _put(sheet, start + 1, column, header)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, TABLE_BORDER
        cell.alignment = HEADER_ALIGN
    sheet.row_dimensions[start + 1].height = 34
    for offset, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            cell = _put(sheet, start + offset, column, value,
                        FMT_DATETIME if isinstance(value, datetime) else None)
            cell.border = TABLE_BORDER
            if offset % 2 == 0:
                cell.fill = ZEBRA_FILL
        sheet.row_dimensions[start + offset].height = 30
    return start + len(rows) + 3


def _summary(sheet, videos, stats, run_log, moment):
    for column, width in enumerate((28, 24, 18, 24, 20, 20, 16, 20), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    row = _table(sheet, 1, 'Tổng quan', ['Chỉ số', 'Giá trị'], [
        ['Thời gian chạy', run_log.get('finished_at', moment).replace(tzinfo=None)],
        ['Số kênh', run_log.get('channel_count', len(stats))],
        ['Video quét', sum(item['scanned'] for item in stats)],
        ['Video lấy mới', sum(not video.taken_on for video in videos)],
        ['Video lấy lại', sum(bool(video.taken_on) for video in videos)],
        ['Tổng lượt xem', format_views(sum(video.views or 0 for video in videos))]])
    channel_rows = [[item.get('name') or item['url'], format_views(item.get('followers')),
                     format_views(item['scanned']), format_views(item['recent']), format_views(item.get('median')),
                     format_views(item.get('threshold')), format_views(item['selected']),
                     format_views(item.get('highest_views'))] for item in stats]
    header_row = row + 1
    row = _table(sheet, row, 'Kết quả từng kênh',
                 ['Kênh', 'Follower', 'Video quét', 'Video trong N ngày', 'View trung vị',
                  'Ngưỡng lấy', 'Video lấy', 'View cao nhất'], channel_rows)
    sheet.auto_filter.ref = f'A{header_row}:H{header_row + len(stats)}'
    top = sorted(videos, key=lambda video: -(video.views or 0))[:10]
    _table(sheet, row, 'Top 10 lượt xem', ['STT', 'Kênh', 'Lượt xem', 'Link bài'],
           [[index, video.page_name or video.page_url, format_views(video.views), video.url]
            for index, video in enumerate(top, 1)])
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str) and cell.value.startswith(('https://', 'http://')):
                cell.hyperlink = cell.value
                cell.font = Font(name='Calibri', color='0563C1', underline='single', size=11)
    sheet.freeze_panes = 'B3'


def _log(sheet, scan_lines, notes, run_log):
    sheet.column_dimensions['A'].width = 42
    sheet.column_dimensions['B'].width = 100
    params = run_log.get('params', {})
    rows = [['Thời gian bắt đầu', run_log.get('started_at')],
            ['Thời gian kết thúc', run_log.get('finished_at')],
            ['Số ngày quét', params.get('days')],
            ['Lượt xem tối thiểu', format_views(params.get('min_views'))],
            ['Hệ số viral', params.get('viral_factor')],
            ['Tối đa video/kênh', params.get('max_per_channel')],
            ['Loại video', ', '.join({'video': 'Video thường', 'reel': 'Reel'}.get(kind, kind)
                                   for kind in sorted(params.get('formats', [])))],
            ['Model chép lời', 'faster-whisper small (tiếng Việt)'],
            ['AI cải thiện lời thoại', run_log.get('llm') or 'Không bật']]
    rows = [[name, value.replace(tzinfo=None) if isinstance(value, datetime) else value]
            for name, value in rows]
    row = _table(sheet, 1, 'Tham số lượt chạy', ['Tham số', 'Giá trị'], rows)
    row = _table(sheet, row, 'Kết quả quét từng kênh', ['STT', 'Nội dung'],
                 [[index, line] for index, line in enumerate(scan_lines, 1)])
    first_note = row + 2
    _table(sheet, row, 'Ghi chú lượt chạy', ['STT', 'Nội dung'],
           [[index, note if note.startswith(('Bỏ qua', 'LỖI:')) else 'LỖI: ' + note]
            for index, note in enumerate(notes, 1)])
    for offset, note in enumerate(notes):
        sheet.cell(first_note + offset, 2).font = Font(name='Calibri', size=11,
            color='7F7F7F' if note.startswith('Bỏ qua') else 'C0392B')
    sheet.freeze_panes = 'B3'


def export(videos: list[Video], scan_lines: list[str] | None = None,
           notes: list[str] | None = None, output_dir: Path | None = None, *,
           channel_stats: list[dict] | None = None, run_log: dict | None = None) -> Path:
    directory = Path(output_dir) if output_dir is not None else config.OUTPUT_DIR
    directory.mkdir(parents=True, exist_ok=True)
    moment = datetime.now(VN_TZ)
    number = 1
    while True:
        suffix = '' if number == 1 else f' (lần {number})'
        path = directory / f'{moment:%Y-%m-%d}{suffix}.xlsx'
        try:
            handle = path.open('xb')
            break
        except FileExistsError:
            number += 1
    workbook = Workbook()
    try:
        ordered = sorted(videos, key=lambda video: (bool(video.taken_on), -(video.views or 0)))
        detail = workbook.active
        detail.title = SHEET_DETAIL
        _details(detail, ordered)
        _summary(workbook.create_sheet(SHEET_SUMMARY), ordered, channel_stats or [], run_log or {}, moment)
        _log(workbook.create_sheet(SHEET_LOG), scan_lines or [], notes or [], run_log or {})
        for sheet, colour in zip(workbook, ('287D8E', '55A58B', '8295A3')):
            sheet.sheet_properties.tabColor = colour
            sheet.page_setup.orientation = 'landscape'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if sheet is detail else sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.print_options.horizontalCentered = True
            sheet.print_area = sheet.dimensions
        with handle:
            workbook.save(handle)
    except BaseException:
        handle.close()
        path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return path


def read_summary(path: str | Path) -> dict:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError('Không đọc được tệp Excel nội dung.') from exc
    try:
        modern = SHEET_DETAIL in workbook.sheetnames
        name = SHEET_DETAIL if modern else 'Nội dung'
        if name not in workbook.sheetnames:
            raise ValueError('Tệp Excel không đúng định dạng nội dung.')
        rows = workbook[name].iter_rows(min_row=2 if modern else 1, values_only=True)
        header = list(next(rows, ()))
        required = HEADERS if modern else LEGACY_HEADERS
        if any(name not in header for name in required):
            raise ValueError('Tệp Excel thiếu cột nội dung bắt buộc.')
        indices = {name: header.index(name) for name in required}
        entries, channels = [], Counter()
        new = taken = total_views = 0
        for row in rows:
            if not any(value is not None for value in row):
                continue
            def value(name):
                index = indices[name]
                return row[index] if index < len(row) else None
            views_text = str(value('Lượt xem') or '')
            views = parse_count(views_text)
            status = str(value('Tình trạng') or '')
            channel = str(value('Kênh') or '')
            channels[channel] += 1
            new += status.strip().casefold() == 'mới'
            taken += status.startswith('Đã lấy')
            total_views += views or 0
            posted = value('Ngày đăng')
            entries.append({'stt': value('STT'), 'topic': str(value('Chủ đề') or ''),
                            'transcript': str(value('Nội dung thoại') or ''), 'channel': channel,
                            'views': views, 'views_text': views_text,
                            'posted': posted.strftime('%d/%m/%Y') if isinstance(posted, datetime) else str(posted or ''),
                            'url': str(value('Link bài' if modern else 'Link video mẫu') or ''), 'status': status})
        scan_lines, notes, errors = [], [], []
        log_name = SHEET_LOG if modern else 'Tình hình quét'
        if log_name in workbook.sheetnames:
            section = None
            for row in workbook[log_name].iter_rows(values_only=True):
                if modern:
                    if row[0] in ('Kết quả quét từng kênh', 'Ghi chú lượt chạy'):
                        section = row[0]
                        continue
                    if section is None or not isinstance(row[0], int):
                        continue
                    line = str(row[1] or '').strip()
                else:
                    line = str(row[0] or '').strip()
                    if line == 'Kết quả quét từng kênh':
                        continue
                if not line:
                    continue
                if line.startswith('Bỏ qua'):
                    notes.append(line)
                elif line.startswith('LỖI:'):
                    errors.append(line)
                else:
                    scan_lines.append(line)
        return {'total': len(entries), 'new': new, 'taken_before': taken, 'views_total': total_views,
                'channels': [{'name': name, 'count': count} for name, count in channels.most_common()],
                'top': sorted(entries, key=lambda entry: -(entry['views'] or 0))[:10],
                'scan_lines': scan_lines, 'notes': notes, 'errors': errors}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Tệp Excel không đúng định dạng nội dung.') from exc
    finally:
        workbook.close()
